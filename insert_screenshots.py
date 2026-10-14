"""把 demo 截圖插進簡報對應的 Demo 待補 位置，用「等比縮放置中」避免變形。"""
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

PPTX_PATH = "/Users/ian/Desktop/ian-agent/bz-demo/電量Yota顧 提案簡報 v2.pptx"
SHOT_DIR = "/Users/ian/Desktop/ian-agent/bz-demo/demo影片截圖"

# (slide_index(0-based), 目標框 shape_id, 圖片檔名, 要清掉文字的 placeholder shape_id 們)
JOBS = [
    (0, 11, "02_對話診斷_開場快速回覆_slide1LINE對話_slide7用聊天取代表單.png", [13]),   # slide1 LINE對話
    (0, 23, "01_電車月曆_slide1電車月曆_slide1官網預演.png", [25]),                     # slide1 電車月曆
    (6, 15, "02_對話診斷_開場快速回覆_slide1LINE對話_slide7用聊天取代表單.png", [17]),   # slide7 01 用聊天取代表單
    (6, 24, "03_對話診斷_聽懂不規則描述_slide7聽懂不規則的描述.png", [26]),              # slide7 02 聽懂不規則的描述
    (6, 33, "04_適配度結果_slide7算出油轉電適配度.png", [35]),                          # slide7 03 算出油轉電適配度
    (8, 39, "05_焦慮雷達_車主反饋卡_slide9車主反饋卡.png", [41]),                       # slide9 車主反饋卡
    (9, 18, "06_業務摘要卡_iPad_slide10業務摘要卡.png", [20]),                          # slide10 業務摘要卡
]


def fit_box(img_w, img_h, box_l, box_t, box_w, box_h):
    scale = min(box_w / img_w, box_h / img_h)
    new_w, new_h = img_w * scale, img_h * scale
    left = box_l + (box_w - new_w) / 2
    top = box_t + (box_h - new_h) / 2
    return int(left), int(top), int(new_w), int(new_h)


def main():
    prs = Presentation(PPTX_PATH)
    for slide_idx, target_id, filename, clear_ids in JOBS:
        slide = prs.slides[slide_idx]
        shapes_by_id = {sh.shape_id: sh for sh in slide.shapes}
        target = shapes_by_id[target_id]
        box_l, box_t, box_w, box_h = target.left, target.top, target.width, target.height

        img_path = f"{SHOT_DIR}/{filename}"
        with Image.open(img_path) as im:
            img_w_px, img_h_px = im.size

        left, top, width, height = fit_box(img_w_px, img_h_px, box_l, box_t, box_w, box_h)
        slide.shapes.add_picture(img_path, Emu(left), Emu(top), width=Emu(width), height=Emu(height))

        for cid in clear_ids:
            sh = shapes_by_id.get(cid)
            if sh is not None:
                sh._element.getparent().remove(sh._element)

        print(f"slide{slide_idx+1}: 放入 {filename} 到 shape#{target_id} 框（{box_w}x{box_h} EMU），"
              f"圖片等比縮成 {width}x{height}，清掉 placeholder {clear_ids}")

    prs.save(PPTX_PATH)
    print("已存檔。")


if __name__ == "__main__":
    main()
